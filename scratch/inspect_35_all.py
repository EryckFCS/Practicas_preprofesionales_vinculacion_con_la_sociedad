import os
import openpyxl
from src.config import RESULTADOS_DIR

excel_path = os.path.join(RESULTADOS_DIR, "35 IND-REPRODUCC ANIMAL 2025.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (dimensions: {sheet.dimensions}) ---")
    for r_idx in range(1, 10):
        row_vals = [sheet.cell(r_idx, col).value for col in range(1, 8)]
        if any(v is not None for v in row_vals):
            print(f"Row {r_idx}: {row_vals}")
