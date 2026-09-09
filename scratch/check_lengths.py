import docx

doc_orig = docx.Document("/home/erick-fcs/Descargas/Practicas_preprofesionales/INFORMES 24-46 ERICK/35. INFORME SG REPROD. ANIMAL 2025.docx")
doc_upd = docx.Document("/home/erick-fcs/Descargas/Practicas_preprofesionales/processed/posgrado/35. INFORME SG REPROD. ANIMAL 2025.docx")

print("Original length:", len(doc_orig.paragraphs))
print("Updated length:", len(doc_upd.paragraphs))

print("\nOriginal 'total' matches:")
for idx, p in enumerate(doc_orig.paragraphs):
    if "total" in p.text.lower():
        print(f"  {idx}: {repr(p.text)}")

print("\nUpdated 'total' matches:")
for idx, p in enumerate(doc_upd.paragraphs):
    if "total" in p.text.lower():
        print(f"  {idx}: {repr(p.text)}")
