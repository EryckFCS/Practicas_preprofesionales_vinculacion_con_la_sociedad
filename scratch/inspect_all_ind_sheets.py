import os
import openpyxl
from src.config import CAREER_FILES, RESULTADOS_DIR

for rid in [24, 35]:
    career = CAREER_FILES[rid]
    excel_path = os.path.join(RESULTADOS_DIR, career["excel"])
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ind_sheet = [name for name in wb.sheetnames if "indicador" in name.lower()][0]
    sheet = wb[ind_sheet]
    print(f"\n=== Report {rid}: {career['excel']} ===")
    for r in range(1, 10):
        print(f"Row {r}: {[sheet.cell(r, c).value for c in range(1, 10)]}")
