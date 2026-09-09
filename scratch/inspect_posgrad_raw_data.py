import os
import openpyxl
from src.config import CAREER_FILES, RESULTADOS_DIR

for rid in range(35, 47):
    career = CAREER_FILES[rid]
    excel_path = os.path.join(RESULTADOS_DIR, career["excel"])
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # Let's inspect DataG sheet
    data_g = wb['DataG']
    reg_consulta = data_g.cell(1, 2).value
    total_encuesta = data_g.cell(2, 2).value
    
    # Let's inspect Demográfico sheet
    dem = wb['Demográfico']
    dem_row2 = [dem.cell(2, c).value for c in range(1, 5)]
    
    # Check if there are other sheets or cell values that could indicate population
    print(f"Report {rid}: {career['excel']}")
    print(f"  DataG: registros consulta={reg_consulta}, total encuesta={total_encuesta}")
    print(f"  Demográfico Row 2: {dem_row2}")
