import docx

doc_orig = docx.Document("/home/erick-fcs/Descargas/Practicas_preprofesionales/INFORMES 24-46 ERICK/35. INFORME SG REPROD. ANIMAL 2025.docx")
doc_upd = docx.Document("/home/erick-fcs/Descargas/Practicas_preprofesionales/processed/posgrado/35. INFORME SG REPROD. ANIMAL 2025.docx")

print("=== ORIGINAL ===")
print("p54:", repr(doc_orig.paragraphs[54].text))
print("p112:", repr(doc_orig.paragraphs[112].text))

print("=== UPDATED ===")
print("p54:", repr(doc_upd.paragraphs[54].text))
print("p112:", repr(doc_upd.paragraphs[112].text))
