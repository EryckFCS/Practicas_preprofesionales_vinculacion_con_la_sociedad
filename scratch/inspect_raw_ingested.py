import openpyxl
wb = openpyxl.load_workbook('RESULTADOS-24-46 ERICK/35 IND-REPRODUCC ANIMAL 2025.xlsx', data_only=False)
sheet = wb['Indicadores ']
print("B4 formula/val:", repr(sheet.cell(row=4, column=2).value))
print("C4 formula/val:", repr(sheet.cell(row=4, column=3).value))
