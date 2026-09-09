import docx
import os
import csv
from src.config import CAREER_FILES
from src.extractor import extract_all_data

reports = range(35, 47)
diagnostics = []

print("=== STARTING BATCH DIAGNOSTIC ON REPORTS 35-46 ===")

for r_id in reports:
    career = CAREER_FILES[r_id]
    word_name = career["word"]
    doc_path = f"processed/posgrado/{word_name}"
    
    if not os.path.exists(doc_path):
        print(f"Report {r_id}: Word file not found at {doc_path}")
        continue
    
    print(f"Diagnosing Report {r_id} ({career['open_sheet']})...")
    doc = docx.Document(doc_path)
    data = extract_all_data(r_id)
    
    # 1. Gather basic stats
    n_graduados = data.get("metadata", {}).get("muestra", 0)
    n_empleadores = data.get("employer", {}).get("total_employers", 0)
    if n_empleadores == 0:
        emp_b1 = data.get("employer", {}).get("b1", {})
        n_empleadores = sum(v.get("count", 0) for v in emp_b1.values())
        
    # 2. Check for placeholders in paragraphs
    placeholders_found = []
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        
        # Check standard placeholders
        if "……" in text or "….." in text or "Conocimientos en …" in text or "Temática 1" in text or "Temática 2" in text:
            placeholders_found.append(f"Para {idx}: '{text[:50]}...'")
            
    # 3. Check for all-zero tables
    all_zero_tables = []
    for t_idx, table in enumerate(doc.tables):
        has_numeric = False
        all_zeros = True
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            for val_str in row_cells:
                clean = val_str.replace('%', '').replace(',', '.').strip()
                try:
                    val = float(clean)
                    has_numeric = True
                    if val != 0.0:
                        all_zeros = False
                except ValueError:
                    pass
        if has_numeric and all_zeros:
            all_zero_tables.append(f"Table {t_idx}")
            
    # 4. Determine final status
    status = "OK"
    issues = []
    if placeholders_found:
        status = "PLACEHOLDERS PENDING"
        issues.append(f"{len(placeholders_found)} placeholders")
    if all_zero_tables:
        issues.append(f"All-zero tables ({', '.join(all_zero_tables)})")
        
    issue_desc = "; ".join(issues) if issues else "Ninguno (Totalmente limpio)"
    
    diagnostics.append({
        "report_id": r_id,
        "carrera": career["open_sheet"],
        "graduados_muestra": n_graduados,
        "empleadores_muestra": n_empleadores,
        "placeholders_detectados": len(placeholders_found),
        "tablas_en_cero": len(all_zero_tables),
        "estado_final": status,
        "observaciones": issue_desc
    })

csv_path = "/home/erick-fcs/Descargas/Practicas_preprofesionales/scratch/diagnostico_lote_posgrado.csv"
with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=["report_id", "carrera", "graduados_muestra", "empleadores_muestra", "placeholders_detectados", "tablas_en_cero", "estado_final", "observaciones"])
    writer.writeheader()
    writer.writerows(diagnostics)

print(f"\nBatch diagnostic complete. CSV written to {csv_path}")
