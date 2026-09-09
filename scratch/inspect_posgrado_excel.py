import os
import openpyxl
from src.config import CAREER_FILES, RESULTADOS_DIR

for rid in range(35, 47):
    career = CAREER_FILES[rid]
    excel_path = os.path.join(RESULTADOS_DIR, career["excel"])
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    print(f"\n=== Report {rid}: {career['excel']} ===")
    print("Sheets:", wb.sheetnames)
    
    # check for Indicators sheet
    ind_sheet = None
    for name in wb.sheetnames:
        if "indicador" in name.lower():
            ind_sheet = wb[name]
            break
    
    if ind_sheet:
        print("Indicadores sheet rows:")
        for r_idx, row in enumerate(ind_sheet.iter_rows(values_only=True), 1):
            row_str = " | ".join(str(val) for val in row if val is not None)
            if "total" in row_str.lower() or "poblaci" in row_str.lower() or "muestra" in row_str.lower():
                print(f"Row {r_idx}: {row_str}")
    else:
        print("No Indicadores sheet found")
